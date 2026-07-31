#!/usr/bin/env python3
"""Generate a deterministic synthetic DOCX covering redaction edge cases."""

from __future__ import annotations

import argparse
import tempfile
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
VT_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
CP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"


def qn(namespace: str, name: str) -> str:
    return f"{{{namespace}}}{name}"


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn(W_NS, "eastAsia"), name)


def add_relationship(xml: bytes, relationship_type: str, target: str) -> bytes:
    root = etree.fromstring(xml)
    used = {
        int(value[3:])
        for value in root.xpath("./pr:Relationship/@Id", namespaces={"pr": REL_NS})
        if value.startswith("rId") and value[3:].isdigit()
    }
    relation_id = next(value for value in range(1, 10_000) if value not in used)
    relationship = etree.SubElement(root, qn(REL_NS, "Relationship"))
    relationship.set("Id", f"rId{relation_id}")
    relationship.set("Type", relationship_type)
    relationship.set("Target", target)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def add_override(xml: bytes, part_name: str, content_type: str) -> bytes:
    root = etree.fromstring(xml)
    existing = root.xpath(
        "./ct:Override[@PartName=$part]",
        namespaces={"ct": CT_NS},
        part=part_name,
    )
    if not existing:
        override = etree.SubElement(root, qn(CT_NS, "Override"))
        override.set("PartName", part_name)
        override.set("ContentType", content_type)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def patch_document_xml(xml: bytes) -> bytes:
    root = etree.fromstring(xml)
    namespaces = {"w": W_NS}

    comment_text = "项目仅供内部使用"
    comment_node = root.xpath(
        ".//w:t[text()=$text]", namespaces=namespaces, text=comment_text
    )[0]
    comment_run = comment_node.getparent()
    parent = comment_run.getparent()
    index = parent.index(comment_run)
    start = etree.Element(qn(W_NS, "commentRangeStart"))
    start.set(qn(W_NS, "id"), "0")
    end = etree.Element(qn(W_NS, "commentRangeEnd"))
    end.set(qn(W_NS, "id"), "0")
    reference_run = etree.Element(qn(W_NS, "r"))
    reference = etree.SubElement(reference_run, qn(W_NS, "commentReference"))
    reference.set(qn(W_NS, "id"), "0")
    parent.insert(index, start)
    parent.insert(index + 2, end)
    parent.insert(index + 3, reference_run)

    footnote_marker = root.xpath(
        ".//w:t[text()='[[FOOTNOTE]]']", namespaces=namespaces
    )[0]
    footnote_run = footnote_marker.getparent()
    footnote_run.remove(footnote_marker)
    footnote_ref = etree.SubElement(footnote_run, qn(W_NS, "footnoteReference"))
    footnote_ref.set(qn(W_NS, "id"), "2")

    tracked_marker = root.xpath(
        ".//w:t[text()='[[TRACKED]]']", namespaces=namespaces
    )[0]
    tracked_run = tracked_marker.getparent()
    tracked_parent = tracked_run.getparent()
    tracked_index = tracked_parent.index(tracked_run)
    tracked_parent.remove(tracked_run)
    deletion = etree.Element(qn(W_NS, "del"))
    deletion.set(qn(W_NS, "id"), "8")
    deletion.set(qn(W_NS, "author"), "Sensitive Reviewer")
    deletion.set(qn(W_NS, "date"), "2026-07-31T12:00:00Z")
    deleted_run = etree.SubElement(deletion, qn(W_NS, "r"))
    deleted_text = etree.SubElement(deleted_run, qn(W_NS, "delText"))
    deleted_text.text = "旧手机号：13900000002"
    tracked_parent.insert(tracked_index, deletion)

    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def comments_xml() -> bytes:
    root = etree.Element(qn(W_NS, "comments"), nsmap={"w": W_NS})
    comment = etree.SubElement(root, qn(W_NS, "comment"))
    comment.set(qn(W_NS, "id"), "0")
    comment.set(qn(W_NS, "author"), "Sensitive Reviewer")
    comment.set(qn(W_NS, "initials"), "SR")
    comment.set(qn(W_NS, "date"), "2026-07-31T12:00:00Z")
    paragraph = etree.SubElement(comment, qn(W_NS, "p"))
    run = etree.SubElement(paragraph, qn(W_NS, "r"))
    text = etree.SubElement(run, qn(W_NS, "t"))
    text.text = "批注联系人：comment@example.com"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def footnotes_xml() -> bytes:
    root = etree.Element(qn(W_NS, "footnotes"), nsmap={"w": W_NS})
    for note_id, note_type in (("-1", "separator"), ("0", "continuationSeparator")):
        note = etree.SubElement(root, qn(W_NS, "footnote"))
        note.set(qn(W_NS, "id"), note_id)
        paragraph = etree.SubElement(note, qn(W_NS, "p"))
        run = etree.SubElement(paragraph, qn(W_NS, "r"))
        etree.SubElement(run, qn(W_NS, note_type))
    note = etree.SubElement(root, qn(W_NS, "footnote"))
    note.set(qn(W_NS, "id"), "2")
    paragraph = etree.SubElement(note, qn(W_NS, "p"))
    run = etree.SubElement(paragraph, qn(W_NS, "r"))
    text = etree.SubElement(run, qn(W_NS, "t"))
    text.text = "脚注邮箱：footnote@example.com"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def custom_properties_xml() -> bytes:
    root = etree.Element(qn(CP_NS, "Properties"), nsmap={None: CP_NS, "vt": VT_NS})
    prop = etree.SubElement(root, qn(CP_NS, "property"))
    prop.set("fmtid", "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}")
    prop.set("pid", "2")
    prop.set("name", "InternalCustomer")
    value = etree.SubElement(prop, qn(VT_NS, "lpwstr"))
    value.text = "星河示例数据研究院"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def patch_package(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}

    entries["word/document.xml"] = patch_document_xml(entries["word/document.xml"])
    entries["word/comments.xml"] = comments_xml()
    entries["word/footnotes.xml"] = footnotes_xml()
    entries["docProps/custom.xml"] = custom_properties_xml()

    relationships = entries["word/_rels/document.xml.rels"]
    relationships = add_relationship(
        relationships,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
        "comments.xml",
    )
    relationships = add_relationship(
        relationships,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        "footnotes.xml",
    )
    entries["word/_rels/document.xml.rels"] = relationships

    package_rels = entries["_rels/.rels"]
    package_rels = add_relationship(
        package_rels,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties",
        "docProps/custom.xml",
    )
    entries["_rels/.rels"] = package_rels

    content_types = entries["[Content_Types].xml"]
    content_types = add_override(
        content_types,
        "/word/comments.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    content_types = add_override(
        content_types,
        "/word/footnotes.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    )
    content_types = add_override(
        content_types,
        "/docProps/custom.xml",
        "application/vnd.openxmlformats-officedocument.custom-properties+xml",
    )
    entries["[Content_Types].xml"] = content_types

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=path.parent) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name in sorted(entries):
                info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o100644 << 16
                target.writestr(info, entries[name])
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def build(path: Path, image: Path | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = document.styles
    styles["Normal"].font.name = "Arial Unicode MS"
    styles["Normal"]._element.get_or_add_rPr().rFonts.set(
        qn(W_NS, "eastAsia"), "Arial Unicode MS"
    )
    styles["Normal"].font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("LlaMask DOCX 脱敏回归样本")
    set_east_asia_font(title_run, "Arial Unicode MS")
    title_run.bold = True
    title_run.font.size = Pt(18)

    if image is not None:
        document.add_paragraph("下方图片包含仅用于回归测试的合成敏感数据：")
        document.add_picture(str(image), width=Inches(6.0))

    paragraph = document.add_paragraph("跨 run 联系电话：")
    run = paragraph.add_run("138")
    run.bold = True
    run = paragraph.add_run("0000")
    run.italic = True
    run = paragraph.add_run("0001")
    run.underline = True
    paragraph.add_run("，此号码只用于合成测试。")

    document.add_paragraph("居民身份证号码：110105198003150020")
    document.add_paragraph("项目仅供内部使用")

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "联系邮箱"
    table.cell(1, 1).text = "table@example.com"

    paragraph = document.add_paragraph("脚注示例：")
    paragraph.add_run("[[FOOTNOTE]]")
    paragraph = document.add_paragraph("修订删除示例：")
    paragraph.add_run("[[TRACKED]]")
    paragraph.add_run("当前公开版本。")

    field_paragraph = document.add_paragraph("隐藏字段代码：")
    field_run = field_paragraph.add_run()
    instr = etree.SubElement(field_run._r, qn(W_NS, "instrText"))
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = ' HYPERLINK "mailto:hidden@example.com" '

    header = section.header.paragraphs[0]
    header.text = "内部合同编号：HT-LLM-2028-00008"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "支持电话：13800000001"

    now = datetime(2026, 7, 31, 12, 0, tzinfo=UTC)
    document.core_properties.author = "Original Author"
    document.core_properties.last_modified_by = "Sensitive Reviewer"
    document.core_properties.created = now
    document.core_properties.modified = now
    document.core_properties.comments = "内部客户：星河示例数据研究院"
    document.save(path)
    patch_package(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "output",
        type=Path,
        nargs="?",
        default=Path("fixtures/docx/comprehensive.docx"),
    )
    parser.add_argument(
        "--image",
        type=Path,
        help="可选的合成 PNG/JPEG，用于生成嵌入图片递归脱敏样本",
    )
    args = parser.parse_args()
    build(args.output, args.image)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
